"""
PyTest configuration and fixtures for the AI Recruitment System.
Provides shared database, Flask test clients, authenticated sessions,
and mock generators for fast, isolated, deterministic test execution.
"""

from __future__ import annotations

import datetime
import io
import os
import shutil
import sys
import tempfile
from collections.abc import Generator
from datetime import timezone
from pathlib import Path

import docx
import pytest
from flask import Flask
from flask.testing import FlaskClient

# Ensure repo root and sub-packages are importable
ROOT = Path(__file__).resolve().parents[1]
for folder in (ROOT, ROOT / "app", ROOT / "src"):
    path_str = str(folder)
    if path_str not in sys.path:
        sys.path.insert(0, path_str)

import config
from app import create_app, create_candidate_app
from app.database import init_db


@pytest.fixture(scope="session", autouse=True)
def configure_test_environment() -> Generator[None, None, None]:
    """Sets environment flags for test isolation."""
    old_env = os.environ.get("FLASK_ENV")
    os.environ["FLASK_ENV"] = "testing"
    os.environ["ARS_TEST_MODE"] = "1"
    config.LOGIN_ENABLED = True
    config.SECRET_KEY = "commercial_grade_test_secret_key_fixed"
    yield
    if old_env is not None:
        os.environ["FLASK_ENV"] = old_env
    else:
        os.environ.pop("FLASK_ENV", None)
    os.environ.pop("ARS_TEST_MODE", None)


@pytest.fixture
def temp_dir() -> Generator[Path, None, None]:
    """Provides a sterile temporary directory cleaned up after test completion."""
    tmp = tempfile.mkdtemp(prefix="ars_test_")
    path = Path(tmp)
    try:
        yield path
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


@pytest.fixture
def test_db(temp_dir: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    """Initializes a dedicated temporary SQLite database with all tables and indexes."""
    db_file = temp_dir / "test_recruitment.db"
    monkeypatch.setattr("app.database.DB_PATH", db_file)
    init_db()
    return db_file


@pytest.fixture
def app(test_db: Path, monkeypatch: pytest.MonkeyPatch) -> Flask:
    """Configured main Flask recruitment application instance."""
    app_instance = create_app()
    app_instance.config["TESTING"] = True
    app_instance.config["SECRET_KEY"] = "commercial_grade_test_secret_key_fixed"
    return app_instance


@pytest.fixture
def candidate_app(test_db: Path, monkeypatch: pytest.MonkeyPatch) -> Flask:
    """Configured isolated candidate remote portal application instance."""
    cand_instance = create_candidate_app()
    cand_instance.config["TESTING"] = True
    cand_instance.config["SECRET_KEY"] = "commercial_grade_test_secret_key_fixed"
    return cand_instance


@pytest.fixture
def client(app: Flask) -> FlaskClient:
    """Unauthenticated HTTP client for the main application."""
    return app.test_client()


@pytest.fixture
def candidate_client(candidate_app: Flask) -> FlaskClient:
    """HTTP client for the remote candidate portal."""
    return candidate_app.test_client()


@pytest.fixture
def auth_client(app: Flask) -> FlaskClient:
    """Authenticated HR Administrator test client with active session and valid CSRF token."""
    c = app.test_client()
    with c.session_transaction() as sess:
        sess["logged_in"] = True
        sess["desktop_session"] = True
        sess["_csrf_token"] = "valid_test_csrf_token"
    c.environ_base["HTTP_X_CSRF_TOKEN"] = "valid_test_csrf_token"
    return c


@pytest.fixture
def sample_pdf_bytes() -> bytes:
    """Generates minimal valid PDF bytes with standard header and trailer."""
    return (
        b"%PDF-1.4\n"
        b"1 0 obj<</Type/Catalog/Pages 2 0 R>>endobj\n"
        b"2 0 obj<</Type/Pages/Count 1/Kids[3 0 R]>>endobj\n"
        b"3 0 obj<</Type/Page/MediaBox[0 0 612 792]/Parent 2 0 R/Resources<<>>>>endobj\n"
        b"xref\n0 4\n0000000000 65535 f \n0000000009 00000 n \n0000000052 00000 n \n0000000101 00000 n \n"
        b"trailer<</Size 4/Root 1 0 R>>\nstartxref\n178\n%%EOF\n"
    )


@pytest.fixture
def sample_docx_bytes() -> bytes:
    """Generates valid DOCX document bytes containing structured resume sections."""
    doc = docx.Document()
    doc.add_heading("Jane Doe - Senior Python Developer", level=1)
    doc.add_paragraph("Email: jane.doe@example.com | Phone: +1 555-0199")
    doc.add_heading("Professional Summary", level=2)
    doc.add_paragraph(
        "Senior Software Engineer with 7 years of hands-on experience designing, "
        "architecting, and deploying distributed backend microservices using Python, "
        "FastAPI, Flask, Docker, Kubernetes, and PostgreSQL."
    )
    doc.add_heading("Technical Skills", level=2)
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Languages"
    table.cell(0, 1).text = "Python, TypeScript, Go, SQL"
    table.cell(1, 0).text = "Frameworks & Cloud"
    table.cell(1, 1).text = "Flask, Django, Docker, AWS, Kubernetes"

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def sample_token_payload() -> dict:
    """Returns valid un-used interview access token metadata."""
    return {
        "token": "T_SAMPLE_TEST_TOKEN",
        "used": 0,
        "candidate_name": "Jane Doe",
        "source_file": "jane_doe_resume.pdf",
        "job_title": "Senior Python Developer",
        "score": 92.5,
        "rank": 1,
        "created_at": datetime.datetime.now(timezone.utc).isoformat(),
    }
