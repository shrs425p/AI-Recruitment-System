import io
import sys

from fpdf import FPDF
from PIL import Image

from src.pdf_to_txt import clean_text, extract_direct_text, extract_docx_text, process_file
from src.security import encrypt_bytes


def test_no_fitz_or_pymupdf_imported():
    """Verify that PyMuPDF (fitz) is not imported by the PDF processing module."""
    import src.pdf_to_txt
    assert "fitz" not in sys.modules or sys.modules.get("fitz") is None or "src.pdf_to_txt" not in getattr(sys.modules.get("fitz", None), "__file__", "")
    # Check module source code does not contain 'import fitz' or 'import pymupdf'
    import inspect
    source = inspect.getsource(src.pdf_to_txt)
    assert "import fitz" not in source
    assert "import pymupdf" not in source


def test_extract_direct_text_normal_pdf(tmp_path):
    """Test standard digital PDF text extraction with pypdf."""
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=14)
    pdf.cell(text="Candidate: Alex Morgan")
    pdf.ln(10)
    pdf.set_font("Helvetica", size=11)
    pdf.cell(text="Skills: Python, Django, PostgreSQL, Docker, AWS")
    pdf_bytes = bytes(pdf.output())

    pdf_file = tmp_path / "alex_resume.pdf"
    pdf_file.write_bytes(pdf_bytes)

    extracted = extract_direct_text(str(pdf_file))
    assert "Alex Morgan" in extracted
    assert "Python, Django, PostgreSQL" in extracted


def test_extract_direct_text_multi_page(tmp_path):
    """Test extracting text across multiple pages."""
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    pdf.cell(text="Page 1 Content - Education: BS Computer Science")
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    pdf.cell(text="Page 2 Content - Work Experience: Senior Software Engineer")
    pdf_bytes = bytes(pdf.output())

    pdf_file = tmp_path / "multipage.pdf"
    pdf_file.write_bytes(pdf_bytes)

    extracted = extract_direct_text(pdf_file)
    assert "Page 1 Content" in extracted
    assert "Page 2 Content" in extracted


def test_extract_direct_text_edge_cases(tmp_path):
    """Test edge cases: nonexistent file, corrupt bytes, encrypted bytes."""
    # Nonexistent file
    assert extract_direct_text(str(tmp_path / "nonexistent.pdf")) == ""

    # Corrupt / invalid bytes
    assert extract_direct_text(b"not a valid pdf header") == ""

    # Encrypted file at rest
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    pdf.cell(text="Secret Candidate Profile")
    raw_pdf = bytes(pdf.output())
    enc_bytes = encrypt_bytes(raw_pdf)

    enc_file = tmp_path / "encrypted.pdf"
    enc_file.write_bytes(enc_bytes)

    assert "Secret Candidate Profile" in extract_direct_text(enc_file)
    assert "Secret Candidate Profile" in extract_direct_text(enc_bytes)


def test_clean_text_normalisation():
    """Test text cleaning routines."""
    dirty = "Line 1\x00\r\n\r\n\r\n\r\nLine 2    with   extra   spaces\r\n"
    cleaned = clean_text(dirty)
    assert "\x00" not in cleaned
    assert "\r" not in cleaned
    assert "Line 1\n\nLine 2 with extra spaces" == cleaned


def test_process_file_with_image_pdf(tmp_path):
    """Test OCR fallback handling on image-based PDFs."""
    # Create an image-based PDF
    img = Image.new("RGB", (200, 100), color="white")
    pdf_buf = io.BytesIO()
    img.save(pdf_buf, format="PDF")
    image_pdf = tmp_path / "scanned_doc.pdf"
    image_pdf.write_bytes(pdf_buf.getvalue())

    out_dir = tmp_path / "txt_output"
    # Even if OCR finds nothing on a blank white image, process_file should not crash
    result = process_file(image_pdf, out_dir)
    assert isinstance(result, bool)


def test_extract_docx_resume(tmp_path):
    """Test Microsoft Word (.docx) resume text extraction and encryption-at-rest handling."""
    import docx

    doc = docx.Document()
    doc.add_heading("Candidate: Priya Sharma", level=1)
    doc.add_paragraph("Summary: Experienced Cloud Infrastructure and DevOps Engineer.")
    doc.add_paragraph("Skills: Kubernetes, Terraform, AWS, Docker, Python")

    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Employer"
    table.rows[0].cells[1].text = "Acme Cloud Corp"
    table.rows[1].cells[0].text = "Role"
    table.rows[1].cells[1].text = "Lead Site Reliability Engineer"

    docx_buf = io.BytesIO()
    doc.save(docx_buf)
    docx_bytes = docx_buf.getvalue()

    # Plain file test
    docx_file = tmp_path / "priya_resume.docx"
    docx_file.write_bytes(docx_bytes)

    extracted = extract_docx_text(docx_file)
    assert "Priya Sharma" in extracted
    assert "Kubernetes, Terraform, AWS" in extracted
    assert "Lead Site Reliability Engineer" in extracted

    # Encrypted file test
    enc_file = tmp_path / "priya_encrypted.docx"
    enc_file.write_bytes(encrypt_bytes(docx_bytes))
    extracted_enc = extract_docx_text(enc_file)
    assert "Priya Sharma" in extracted_enc

    # End-to-end process_file test
    out_dir = tmp_path / "out_txt"
    success = process_file(docx_file, out_dir)
    assert success is True
    target_txt = out_dir / "priya_resume.txt"
    assert target_txt.exists()
    assert "Priya Sharma" in target_txt.read_text(encoding="utf-8")
