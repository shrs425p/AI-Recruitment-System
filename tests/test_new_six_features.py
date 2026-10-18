import json
from pathlib import Path

import config as cfg
from app import create_app
from app.database import (
    create_hr_notification,
    get_all_candidate_notes,
    get_candidate_note,
    get_unread_notification_count,
    init_db,
    save_candidate_note,
)
from app.routes.interview import get_interview_plan_for_role


def test_feature1_docx_resume_support(tmp_path):
    import docx

    from app.routes.upload import ALLOWED_UPLOAD_EXTENSIONS
    from src.pdf_to_txt import extract_docx_text

    assert ".docx" in ALLOWED_UPLOAD_EXTENSIONS

    docx_path = tmp_path / "jane_doe_resume.docx"
    doc = docx.Document()
    doc.add_heading("Jane Doe - Senior Full Stack Engineer", 0)
    doc.add_paragraph("Summary: Experienced engineer in Python, React, and PostgreSQL.")
    table = doc.add_table(rows=1, cols=2)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Skill"
    hdr_cells[1].text = "Level"
    row_cells = table.add_row().cells
    row_cells[0].text = "Python"
    row_cells[1].text = "Expert"
    doc.save(docx_path)

    extracted = extract_docx_text(docx_path)
    assert "Jane Doe" in extracted
    assert "Senior Full Stack Engineer" in extracted
    assert "Python" in extracted
    assert "Level" in extracted


def test_feature2_export_ranking_csv(tmp_path, monkeypatch):
    test_db = tmp_path / "test_csv_export.db"
    monkeypatch.setattr("app.database.DB_PATH", test_db)
    init_db()

    app = create_app()
    client = app.test_client()

    with client.session_transaction() as sess:
        sess["logged_in"] = True
        sess["desktop_session"] = True
        sess["_csrf_token"] = "csrf-test-token"
    client.environ_base["HTTP_X_CSRF_TOKEN"] = "csrf-test-token"

    ranking_dir = tmp_path / "output" / "ranking"
    ranking_dir.mkdir(parents=True, exist_ok=True)
    monkeypatch.setattr("app.routes.ranking.OUTPUT_FOLDER", tmp_path / "output")

    ranking_file = ranking_dir / "ranking_scores_test.json"
    ranking_payload = {
        "job_description": {"job_title": "Backend Lead"},
        "ranked_candidates": [
            {
                "rank": 1,
                "candidate_name": "Alice Johnson",
                "_source_file": "alice_resume",
                "total_score": 92.5,
                "skills_score": 95,
                "experience_score": 90,
                "strengths": ["FastAPI", "PostgreSQL"],
                "hire_recommendation": "Strongly Recommend",
            }
        ]
    }
    with open(ranking_file, "w", encoding="utf-8") as f:
        json.dump(ranking_payload, f)

    resp = client.get("/api/ranking/export/csv")
    assert resp.status_code == 200
    assert resp.headers["Content-Type"].startswith("text/csv")
    csv_text = resp.data.decode("utf-8")
    assert "Rank,Name,Email,Total Score,Skills Match,Experience Score,Strengths,Hire Recommendation" in csv_text
    assert "Alice Johnson" in csv_text
    assert "92.5" in csv_text
    assert "Strongly Recommend" in csv_text


def test_feature3_configurable_interview_length(tmp_path, monkeypatch):
    test_db = tmp_path / "test_interview_cfg.db"
    monkeypatch.setattr("app.database.DB_PATH", test_db)
    init_db()

    # 1. Test custom plan generation
    plan_tech = get_interview_plan_for_role("Software Architect", "IT", tech_count=5, beh_count=3)
    assert len(plan_tech) == 8
    tech_qs = [p for p in plan_tech if p[0] == "TECHNICAL"]
    beh_qs = [p for p in plan_tech if p[0] == "BEHAVIORAL"]
    assert len(tech_qs) == 5
    assert len(beh_qs) == 3

    plan_non_tech = get_interview_plan_for_role("HR Specialist", "Human Resources", tech_count=4, beh_count=1)
    assert len(plan_non_tech) == 5
    domain_qs = [p for p in plan_non_tech if p[0] == "DOMAIN EXPERTISE"]
    beh_qs_2 = [p for p in plan_non_tech if p[0] == "BEHAVIORAL"]
    assert len(domain_qs) == 4
    assert len(beh_qs_2) == 1

    # 2. Test API update of interview settings
    app = create_app()
    client = app.test_client()

    with client.session_transaction() as sess:
        sess["logged_in"] = True
        sess["desktop_session"] = True
        sess["_csrf_token"] = "csrf-test-token"
    client.environ_base["HTTP_X_CSRF_TOKEN"] = "csrf-test-token"

    resp = client.post(
        "/api/settings/interview",
        json={
            "tech_questions": 6,
            "beh_questions": 3,
            "time_limit": 180,
        }
    )
    assert resp.status_code == 200
    data = resp.get_json()
    assert data["success"] is True
    assert data["interview_tech_questions"] == 6
    assert data["interview_beh_questions"] == 3
    assert data["interview_time_limit"] == 180
    assert cfg.INTERVIEW_TECH_QUESTIONS == 6
    assert cfg.INTERVIEW_BEH_QUESTIONS == 3
    assert cfg.INTERVIEW_TIME_LIMIT_SECONDS == 180

    # Reset defaults
    cfg.INTERVIEW_TECH_QUESTIONS = 3
    cfg.INTERVIEW_BEH_QUESTIONS = 2
    cfg.INTERVIEW_TIME_LIMIT_SECONDS = 120


def test_feature4_candidate_notes_and_tags(tmp_path, monkeypatch):
    test_db = tmp_path / "test_notes.db"
    monkeypatch.setattr("app.database.DB_PATH", test_db)
    init_db()

    # DB level tests
    save_candidate_note("Bob Smith", note="Prefers remote role", tags=["Follow up", "Strong Fit"], source_file="bob.pdf")
    retrieved = get_candidate_note("Bob Smith")
    assert retrieved is not None
    assert retrieved["note"] == "Prefers remote role"
    assert "Follow up" in retrieved["tags"]
    assert "Strong Fit" in retrieved["tags"]

    all_notes = get_all_candidate_notes()
    assert "Bob Smith" in all_notes

    # API level tests
    app = create_app()
    client = app.test_client()

    with client.session_transaction() as sess:
        sess["logged_in"] = True
        sess["desktop_session"] = True
        sess["_csrf_token"] = "csrf-test-token"
    client.environ_base["HTTP_X_CSRF_TOKEN"] = "csrf-test-token"

    # POST new note and tag
    resp = client.post(
        "/api/candidates/notes",
        json={
            "candidate_name": "Charlie Brown",
            "note": "Great culture match",
            "tags": ["Hold"],
            "source_file": "charlie.pdf"
        }
    )
    assert resp.status_code == 200
    res_data = resp.get_json()
    assert res_data["success"] is True

    # GET notes
    get_res = client.get("/api/candidates/notes?candidate_name=Charlie Brown")
    assert get_res.status_code == 200
    charlie_data = get_res.get_json()
    assert charlie_data["note"] == "Great culture match"
    assert "Hold" in charlie_data["tags"]


def test_feature5_hr_notifications(tmp_path, monkeypatch):
    test_db = tmp_path / "test_notifications.db"
    monkeypatch.setattr("app.database.DB_PATH", test_db)
    init_db()

    app = create_app()
    client = app.test_client()

    with client.session_transaction() as sess:
        sess["logged_in"] = True
        sess["desktop_session"] = True
        sess["_csrf_token"] = "csrf-test-token"
    client.environ_base["HTTP_X_CSRF_TOKEN"] = "csrf-test-token"

    # Initially 0 unread
    assert get_unread_notification_count() == 0

    # Create notification
    notif_id = create_hr_notification(
        title="Interview Completed: Jane Doe",
        message="Jane Doe has completed the interview for Python Developer. Report is ready.",
        candidate_name="Jane Doe",
        job_title="Python Developer",
        notif_type="interview_completed",
        link="/reports"
    )
    assert isinstance(notif_id, int)
    assert get_unread_notification_count() == 1

    # Check unread count API
    count_resp = client.get("/api/notifications/unread-count")
    assert count_resp.status_code == 200
    assert count_resp.get_json()["unread_count"] == 1

    # Check notifications list API
    list_resp = client.get("/api/notifications?unread_only=true")
    assert list_resp.status_code == 200
    notifs = list_resp.get_json()
    assert len(notifs) == 1
    assert notifs[0]["candidate_name"] == "Jane Doe"

    # Mark as read API
    read_resp = client.post(f"/api/notifications/{notif_id}/read")
    assert read_resp.status_code == 200
    assert get_unread_notification_count() == 0


def test_feature6_air_gapped_offline_integrity():
    root = Path(__file__).resolve().parent.parent

    # 1. Verify local Chart.js bundle exists and is healthy (>100KB)
    chart_js = root / "app" / "static" / "js" / "chart.umd.min.js"
    assert chart_js.exists()
    assert chart_js.stat().st_size > 100_000

    # 2. Verify local fonts exist
    fonts_dir = root / "app" / "static" / "fonts"
    assert (fonts_dir / "Inter-Regular.woff2").exists()
    assert (fonts_dir / "JetBrainsMono-Regular.woff2").exists()

    # 3. Verify zero CDN font links or chart.js CDN links in templates
    templates_to_check = [
        root / "app" / "templates" / "base.html",
        root / "app" / "templates" / "ranking.html",
        root / "app" / "templates" / "reports.html",
        root / "app" / "templates" / "candidate_interview.html",
        root / "app" / "templates" / "login.html",
    ]
    for tmpl in templates_to_check:
        content = tmpl.read_text(encoding="utf-8")
        assert "fonts.googleapis.com" not in content, f"CDN found in {tmpl}"
        assert "fonts.gstatic.com" not in content, f"CDN found in {tmpl}"
        assert "cdn.jsdelivr.net" not in content, f"CDN found in {tmpl}"
